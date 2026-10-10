"""Excel import contracts: no model calls, partial writes or unsafe archive content."""
import copy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
import server
from test_ab import ServiceCase
from ticket_import import parse_excel, template_bytes, ImportError, MAX_UPLOAD, MAX_EXPANDED


def workbook(rows, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.append(headers or ['Subject', 'Message', 'Customer'])
    for row in rows:
        sheet.append(row)
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


class ExcelParser(unittest.TestCase):
    def test_template_round_trip_and_english_chinese_headers(self):
        _, rows = parse_excel(template_bytes())
        self.assertEqual(rows[0][1]['customer'], 'Demo user 09')
        _, rows = parse_excel(workbook([['账单问题', '请退款', '演示客户']], ['标题', '正文', '客户']))
        self.assertEqual(rows[0][1]['subject'], '账单问题')
        self.assertEqual(rows[0][1]['channel'], 'Excel')

    def test_bad_rows_have_specific_issues_and_do_not_echo_formula(self):
        _, rows = parse_excel(workbook([['Missing body', None], ['=HYPERLINK("example.invalid")', 'demo'], ['x' * 121, 'demo']]))
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(row[2] for row in rows))
        self.assertEqual(rows[1][1], {})

    def test_bad_headers_empty_oversized_and_extra_rows_are_rejected(self):
        for raw in (b'not a workbook', b'x' * (MAX_UPLOAD + 1), workbook([], ['Subject']),
                    workbook([], ['Subject', 'Message']), workbook([['one', 'two']], ['Subject', 'Title']),
                    workbook([['one', 'two']] * 101)):
            with self.subTest(size=len(raw)), self.assertRaises(ImportError):
                parse_excel(raw)

    def test_sparse_rows_beyond_the_limit_are_not_silently_truncated(self):
        book = load_workbook(BytesIO(workbook([['demo', 'demo']])))
        book.active['A500'] = 'late ticket'
        out = BytesIO(); book.save(out); book.close()
        with self.assertRaises(ImportError):
            parse_excel(out.getvalue())

    def test_xml_entities_macros_links_and_zip_bombs_are_rejected(self):
        unsafe = [
            ('xl/worksheets/sheet1.xml', b'<!DOCTYPE x [<!ENTITY s "secret">]><x>&s;</x>'),
            ('xl/vbaProject.bin', b'fake macro'),
            ('xl/_rels/workbook.xml.rels', b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid" /></Relationships>'),
            ('large.bin', b'0' * (MAX_EXPANDED + 1)),
        ]
        for name, value in unsafe:
            out = BytesIO()
            with ZipFile(out, 'w', ZIP_DEFLATED) as archive:
                archive.writestr(name, value)
            with self.subTest(name=name), self.assertRaises(ImportError):
                parse_excel(out.getvalue())


class ImportContracts(ServiceCase):
    def preview(self, raw, token=None, origin=None):
        response = {}
        environ = {'REQUEST_METHOD': 'POST', 'PATH_INFO': '/api/import/preview',
                   'HTTP_HOST': '127.0.0.1:8780', 'HTTP_X_CSRF_TOKEN': token or server.CSRF_TOKEN,
                   'CONTENT_TYPE': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                   'CONTENT_LENGTH': str(len(raw)), 'wsgi.input': BytesIO(raw)}
        if origin:
            environ['HTTP_ORIGIN'] = origin
        def start(status, headers):
            response['status'] = int(status.split()[0])
        response['body'] = json.loads(b''.join(server.application(environ, start)))
        return response

    def test_preview_no_persistence_and_duplicates_classified(self):
        original = copy.deepcopy(server.tickets)
        result = self.preview(workbook([['Demo request', 'Demo message', 'Demo customer']] * 2))
        self.assertEqual(result['status'], 200)
        self.assertEqual((result['body']['ready'], result['body']['duplicates']), (1, 1))
        self.assertEqual(server.tickets, original)
        self.local_mock.assert_not_called(); self.remote_mock.assert_not_called()

    def test_batch_commit_retry_and_exact_duplicates_never_double_import(self):
        body = {'rows': [{'subject': 'Demo import', 'message': 'Synthetic feedback'}] * 2, 'request_id': str(uuid4())}
        result = self.call('import', body, 201)
        self.assertEqual((result['imported'], result['skipped']), (1, 1))
        self.assertEqual(self.call('import', body)['imported'], 1)
        self.assertEqual(len(server.tickets), 9)
        self.assertEqual(self.call('import', {**body, 'request_id': str(uuid4())}, 201)['imported'], 0)
        self.call('import', {**body, 'rows': [{'subject': 'Changed', 'message': 'different'}]}, 409)
        self.assertEqual(server.tickets[0]['runs']['jev']['status'], 'pending')
        self.local_mock.assert_not_called(); self.remote_mock.assert_not_called()

    def test_invalid_batch_capacity_and_disk_failure_are_atomic(self):
        before = copy.deepcopy(server.tickets)
        body = {'request_id': str(uuid4()), 'rows': [{'subject': 'Valid', 'message': 'Demo'}, {'subject': 'Invalid'}]}
        self.call('import', body, 400)
        self.assertEqual(server.tickets, before)
        body['rows'] = [{'subject': str(i), 'message': 'Demo'} for i in range(100)]
        self.call('import', body, 409)
        self.assertEqual(server.tickets, before)
        body['rows'] = [{'subject': 'Valid', 'message': 'Demo'}]
        with patch.object(server.TicketStore, 'save', side_effect=OSError('disk full')):
            self.call('import', body, 503)
        self.assertEqual(server.tickets, before)

    def test_preview_requires_csrf_and_trusted_origin(self):
        raw = workbook([['Demo', 'Synthetic message']])
        self.assertEqual(self.preview(raw, token='wrong')['status'], 403)
        self.assertEqual(self.preview(raw, origin='https://example.invalid')['status'], 403)

    def test_done_reopen_are_model_specific_and_revision_checked(self):
        result = self.call('resolve', {'id': 'TK-1042', 'expected_revision': 0})['ticket']
        self.assertEqual(result['status'], 'done')
        self.assertEqual(result['runs']['jev']['status'], 'pending')
        self.call('analyze', {'id': 'TK-1042'}, 409)
        self.call('reopen', {'id': 'TK-1042', 'expected_revision': 0}, 409)
        result = self.call('reopen', {'id': 'TK-1042', 'expected_revision': 1})['ticket']
        self.assertEqual(result['status'], 'pending')
