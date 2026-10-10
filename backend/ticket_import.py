"""Bounded, local-only Excel parsing. Uploaded workbooks are never persisted."""
from hashlib import sha256
from io import BytesIO
import json
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

MAX_UPLOAD = 2 * 1024 * 1024
MAX_IMPORT_BODY = 2 * 1024 * 1024
MAX_ROWS = 100
MAX_COLUMNS = 30
MAX_EXPANDED = 8 * 1024 * 1024
FIELDS = {'subject': 120, 'message': 3000, 'customer': 100, 'company': 100, 'channel': 40}
SOURCE_FIELDS = {'source': 24, 'source_id': 256}
ALIASES = {
    'subject': 'subject', 'title': 'subject', 'ticket title': 'subject', '标题': 'subject', '工单标题': 'subject',
    'message': 'message', 'body': 'message', 'description': 'message', 'content': 'message', '正文': 'message', '内容': 'message', '描述': 'message',
    'customer': 'customer', 'name': 'customer', '客户': 'customer', '客户名称': 'customer',
    'company': 'company', 'organization': 'company', '公司': 'company',
    'channel': 'channel', '渠道': 'channel',
}
DEFAULTS = {'customer': 'Imported customer', 'company': 'Imported tickets', 'channel': 'Excel'}


class ImportError(ValueError):
    pass


def normalize_row(row):
    if not isinstance(row, dict) or set(row) - set(FIELDS) - set(SOURCE_FIELDS):
        raise ImportError('Invalid ticket columns.')
    result = {}
    for field, limit in FIELDS.items():
        value = row.get(field, '')
        if not isinstance(value, str):
            raise ImportError(f'{field.title()} must be text.')
        value = value.strip()
        if len(value) > limit:
            raise ImportError(f'{field.title()}: maximum {limit} characters.')
        if any(ord(c) < 32 and c not in '\n\r\t' for c in value):
            raise ImportError(f'{field.title()} contains unsupported characters.')
        if field in ('subject', 'message') and not value:
            raise ImportError(f'{field.title()} is required.')
        result[field] = value or DEFAULTS.get(field, '')
    for field, limit in SOURCE_FIELDS.items():
        value = row.get(field, '')
        if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 for c in value):
            raise ImportError('Invalid source identity.')
        if value:
            result[field] = value
    return result


def row_key(row):
    if row.get('source') and row.get('source_id'):
        return sha256(json.dumps([row['source'], row['source_id']]).encode()).hexdigest()
    return sha256(json.dumps({k: row.get(k, '') for k in FIELDS}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def batch_digest(rows):
    return sha256(json.dumps(rows, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def classify_rows(rows, existing):
    seen = {row_key(t) for t in existing}
    result = []
    for number, row, error in rows:
        if error:
            result.append({'row': number, 'status': 'error', 'issue': error, 'ticket': row})
            continue
        key = row_key(row)
        duplicate = key in seen
        seen.add(key)
        result.append({'row': number, 'status': 'duplicate' if duplicate else 'ready',
                       'issue': 'Duplicate ticket' if duplicate else '', 'ticket': row})
    return result


def parse_excel(raw):
    if not 0 < len(raw) <= MAX_UPLOAD:
        raise ImportError('Choose an .xlsx file under 2 MB.')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 200 or len(names) != len(set(names)):
                raise ImportError('Workbook has too many or duplicate parts.')
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED:
                raise ImportError('Workbook expands beyond the 8 MB limit.')
            for entry in entries:
                path = PurePosixPath(entry.filename)
                lower = entry.filename.lower()
                if path.is_absolute() or '..' in path.parts or '\\' in lower or ':' in lower or entry.flag_bits & 1:
                    raise ImportError('Unsupported workbook archive.')
                if any(part in lower for part in ('vbaproject', 'externallinks/', 'embeddings/')):
                    raise ImportError('Macros, external links and embedded files are not supported.')
                if lower.endswith(('.xml', '.rels')):
                    xml = archive.read(entry)
                    root = fromstring(xml)
                    if lower.startswith('xl/worksheets/'):
                        for row in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row'):
                            if int(row.get('r', '0')) > MAX_ROWS + 1 and len(row):
                                raise ImportError('Use at most 100 rows per import.')
                    if lower == '[content_types].xml' and any('macroenabled' in str(node.attrib).lower() for node in root):
                        raise ImportError('Macro-enabled workbooks are not supported.')
                    if lower.endswith('.rels') and any(node.get('TargetMode') == 'External' for node in root):
                        raise ImportError('External links are not supported.')
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ImportError('Use a workbook with one worksheet.')
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1)
            headers = next(iterator, ())
            mapping = {}
            for index, cell in enumerate(headers):
                if cell.value is None:
                    continue
                if cell.data_type == 'f' or not isinstance(cell.value, str):
                    raise ImportError('Column headers must be plain text.')
                field = ALIASES.get(cell.value.strip().lower())
                if not field:
                    raise ImportError('Unknown column. Use Subject, Message, Customer, Company or Channel.')
                if field in mapping.values():
                    raise ImportError(f'Duplicate {field.title()} column.')
                if index >= MAX_COLUMNS:
                    raise ImportError('Workbook has too many columns.')
                mapping[index] = field
            if not {'subject', 'message'} <= set(mapping.values()):
                raise ImportError('Subject and Message columns are required.')
            rows = []
            for number, cells in enumerate(iterator, 2):
                if not any(cell.value is not None for cell in cells):
                    continue
                if number > MAX_ROWS + 1:
                    raise ImportError('Use at most 100 rows per import.')
                values = {field: '' if cells[index].value is None else cells[index].value for index, field in mapping.items()}
                issue = 'Use plain text; formulas and Excel errors are not supported.' if any(c.data_type in ('f', 'e') for c in cells) else ''
                if any(c.value is not None for i, c in enumerate(cells) if i not in mapping):
                    issue = 'Data appears in a column without a header.'
                if not issue:
                    try:
                        values = normalize_row(values)
                    except ImportError as exc:
                        issue = str(exc)
                # Invalid cell values are not echoed; the row number and issue are enough.
                rows.append((number, values if not issue else {}, issue))
            if not rows:
                raise ImportError('The worksheet has no tickets.')
            return sheet.title, rows
        finally:
            workbook.close()
    except ImportError:
        raise
    except (BadZipFile, DefusedXmlException, ParseError, InvalidFileException, ValueError, KeyError, OSError, TypeError, IndexError, RuntimeError, NotImplementedError):
        raise ImportError('Invalid or unsafe .xlsx workbook.') from None


def template_bytes():
    workbook = Workbook()
    workbook.properties.creator = 'Model Lab'
    sheet = workbook.active
    sheet.title = 'Tickets'
    sheet.append(['Subject', 'Message', 'Customer', 'Company', 'Channel'])
    sheet.append(['Duplicate subscription charge', 'Our demo subscription was charged twice. Please refund the extra payment.',
                  'Demo user 09', 'Example team 09', 'Email'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:E2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='237A67')
    for key, width in {'A': 38, 'B': 80, 'C': 22, 'D': 24, 'E': 16}.items():
        sheet.column_dimensions[key].width = width
    sheet['B2'].alignment = Alignment(wrap_text=True)
    out = BytesIO()
    workbook.save(out)
    workbook.close()
    return out.getvalue()
